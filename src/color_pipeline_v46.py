"""Standalone v46 colour workflow: prepare, detect, calibrate and export.

All stages are normal Python functions. Shared analysis state belongs to one
AnalysisSession; no source rewriting or legacy detector fallback is performed.
"""
from analysis_session_v46 import AnalysisSession
import legend_palette_v46 as palette


def run(argv, legend_runtime, marker_runtime):
    session = AnalysisSession()
    session.values.update(_CLI_ARGV=list(argv), _V46_LEGEND=legend_runtime,
                          _V46_MARKERS=marker_runtime)
    for name in ('filter_palette_noise', 'recover_palette', 'set_palette_names',
                 'lock_palette', 'sync_palette'):
        session.install(getattr(palette, name))
    session.run(_prepare_image_and_palette)
    legend_runtime.prepare(session.values)
    marker_runtime.run(session.values, legend_runtime)
    session.run(_export_results)
    return session.values


def _prepare_image_and_palette():
    global AXIS_COLS, AXIS_ROWS, BG_MASK, COLORS, DARK_BG, FILLED_AREA, GREY_PEAK_DIST_EFF, H, HAS_TICKS, IMG_PATH, LEGEND_BOX, LEGEND_IMG_PATH, NOISE_LAB, NOISE_MATCH_DIST, NO_STEM, OUT_DIR, PLOT_AREA, PLOT_MASK, TICK_XS, USER_LEGEND_BOX, USER_PLOT_AREA, USER_X_LOG, USER_X_MAX, USER_X_MIN, USER_Y_LOG, USER_Y_MAX, USER_Y_MIN, W, _BORDER_V, _BOX_EDGES, _ap, _args, _cfg, _dig, _ex0, _ex1, _extra_exclude, _ey0, _ey1, _gf, _hsv, _lb, _lb_cy, _lx0, _lx1, _ly0, _ly1, _overlap_w, _pa, _pa_h, _pa_w, _pax0, _pax1, _pay0, _pay1, _stem, _ux0, _ux1, _uy0, _uy1, _wide, all_detections, all_results, all_tcaps_out, all_wm_full, img, img_hsv, img_lab, img_rgb
    _ap = _argparse.ArgumentParser()
    _ap.add_argument('image', help='Input chart image')
    _ap.add_argument('output_dir', nargs='?', default=None, help='Output directory')
    _ap.add_argument('--legend', default=None, help='External legend image path')
    _ap.add_argument('--no-stem', action='store_true', help='Skip stem/pixel classification; use raw mask directly for walk')
    _ap.add_argument('--legend-box', default=None, help='User-drawn legend box "x0,y0,x1,y1" (overrides auto-detect)')
    _ap.add_argument('--plot-area', default=None, help='User-drawn plot area "x0,y0,x1,y1" (overrides auto-detect)')
    _ap.add_argument('--x-min', default=None, help='Manual x-axis minimum (data value at plot-area left edge)')
    _ap.add_argument('--x-max', default=None, help='Manual x-axis maximum (data value at plot-area right edge)')
    _ap.add_argument('--y-min', default=None, help='Manual y-axis minimum (data value at plot-area bottom edge)')
    _ap.add_argument('--y-max', default=None, help='Manual y-axis maximum (data value at plot-area top edge)')
    _ap.add_argument('--x-log', action='store_true', help='Treat x-axis as log scale')
    _ap.add_argument('--y-log', action='store_true', help='Treat y-axis as log scale')
    _ap.add_argument('--outlier-filter', action='store_true', help='Run optional post-detection outlier / fake-segment filters')
    _args = _ap.parse_args(_CLI_ARGV)
    USER_LEGEND_BOX = _parse_box(_args.legend_box)
    USER_PLOT_AREA = _parse_box(_args.plot_area)
    USER_X_MIN = _parse_num(_args.x_min)
    USER_X_MAX = _parse_num(_args.x_max)
    USER_Y_MIN = _parse_num(_args.y_min)
    USER_Y_MAX = _parse_num(_args.y_max)
    USER_X_LOG = bool(_args.x_log)
    USER_Y_LOG = bool(_args.y_log)
    if not _args.image:
        print('Usage: python run_A4_auto_v46.py <image_path> [output_dir] [--legend legend_img]')
        sys.exit(1)
    IMG_PATH = _args.image
    LEGEND_IMG_PATH = _args.legend
    NO_STEM = _args.no_stem
    if _args.output_dir:
        OUT_DIR = _args.output_dir
    else:
        _stem = os.path.splitext(os.path.basename(IMG_PATH))[0]
        _stem = re.sub('^pasted_file_', '', _stem)
        _stem = re.sub('_image$', '', _stem)
        OUT_DIR = os.path.join(os.path.dirname(os.path.abspath(IMG_PATH)), f'{_stem}_v46_out')
    os.makedirs(OUT_DIR, exist_ok=True)
    img = cv2.imread(IMG_PATH)
    if img is None:
        print(f'ERROR: cannot read {IMG_PATH}', file=sys.stderr)
        sys.exit(1)
    img_rgb = cv2.cvtColor(img, cv2.COLOR_BGR2RGB)
    img_hsv = cv2.cvtColor(img, cv2.COLOR_BGR2HSV).astype(np.int32)
    img_lab = cv2.cvtColor(img, cv2.COLOR_BGR2Lab).astype(np.float32)
    H, W = img.shape[:2]
    print(f'Image: {W}x{H}  ({IMG_PATH})')
    print('BUILD SIGNATURE: standalone colour v46')
    DARK_BG = _detect_dark_background()
    print(f'Background: {('dark' if DARK_BG else 'light')}')
    _BORDER_V = _border_brightness()
    GREY_PEAK_DIST_EFF = 40 if _BORDER_V >= 248 else 60
    print(f'Border brightness: {_BORDER_V:.0f}  ->  grey peak-distance = {GREY_PEAK_DIST_EFF}')
    BG_MASK = _pcm_background_mask(img_hsv)
    AXIS_ROWS, AXIS_COLS = _detect_axes()
    print(f'Axis rows: {AXIS_ROWS.tolist()},  Axis cols: {AXIS_COLS.tolist()}')
    try:
        _BOX_EDGES = _detect_plot_box_edges()
        print(f'PLOT_BOX: {_BOX_EDGES}')
    except Exception:
        _BOX_EDGES = {}
    FILLED_AREA = _detect_filled_plot_area()
    PLOT_AREA = _compute_plot_area()
    if USER_PLOT_AREA is not None:
        PLOT_AREA = USER_PLOT_AREA
        _ux0, _uy0, _ux1, _uy1 = USER_PLOT_AREA
        AXIS_COLS = np.array([int(_ux0)])
        AXIS_ROWS = np.array([int(_uy1)])
        print(f'Plot area (user-provided): {PLOT_AREA}')
        print(f'Axis rows: {AXIS_ROWS},  Axis cols: {AXIS_COLS}  (from user plot area)')
    else:
        print(f'Plot area: {PLOT_AREA}')
    PLOT_MASK = _plot_area_mask()
    NOISE_LAB = _detect_noise_color()
    print(f'Noise (axis/text/LLOQ) colour Lab: {NOISE_LAB.astype(int).tolist()}')
    NOISE_MATCH_DIST = 45.0
    HAS_TICKS, TICK_XS = _detect_ticks()
    print('\n--- Stage 2: Colour discovery ---')
    COLORS, LEGEND_BOX = _V46_LEGEND.discover(_discover_colors, globals())
    print(f'FINAL_LEGEND_BOX: {(LEGEND_BOX if LEGEND_BOX else 'NONE')}')
    filter_palette_noise()
    print(f'\n--- Stage 3-8: Per-colour processing ---')
    all_results = {}
    all_detections = {}
    all_tcaps_out = {}
    all_wm_full = {}
    _pa = PLOT_AREA
    _cfg = PlotConfig()
    _lb = tuple((int(v) for v in LEGEND_BOX)) if LEGEND_BOX else None
    _pax0, _pay0, _pax1, _pay1 = (int(_pa[0]), int(_pa[1]), int(_pa[2]), int(_pa[3]))
    _extra_exclude = None
    if _lb:
        _lx0, _ly0, _lx1, _ly1 = _lb
        _pa_w = max(1, _pax1 - _pax0)
        _pa_h = max(1, _pay1 - _pay0)
        _overlap_w = max(0, min(_pax1, _lx1) - max(_pax0, _lx0))
        _wide = _overlap_w > 0.5 * _pa_w
        _lb_cy = 0.5 * (_ly0 + _ly1)
        if _wide and _lb_cy < _pay0 + _pa_h * 0.33 and (_ly0 > _pay0 + 2):
            _extra_exclude = (_pax0, _pay0, _pax1, _ly0 - 1)
            print(f'  [excluding title band above legend: y {_pay0}-{_ly0 - 1}]')
        elif _wide and _lb_cy > _pay1 - _pa_h * 0.33 and (_ly1 < _pay1 - 2):
            _extra_exclude = (_pax0, _ly1 + 1, _pax1, _pay1)
            print(f'  [excluding band below legend: y {_ly1 + 1}-{_pay1}]')
    _cfg.plot_area = (_pax0, _pay0, _pax1, _pay1)
    _cfg.legend_box = _lb
    if (_pay0, _pay1) != (int(_pa[1]), int(_pa[3])):
        print(f'  [plot-area trimmed to exclude legend: y {int(_pa[1])}-{int(_pa[3])} -> {_pay0}-{_pay1}]')
    _dig = PlotDigitizer.__new__(PlotDigitizer)
    _dig.cfg = _cfg
    _dig.img = img
    _dig.H, _dig.W = img.shape[:2]
    _dig.rgb = cv2.cvtColor(img, cv2.COLOR_BGR2RGB)
    _dig.lab = cv2.cvtColor(_dig.rgb, cv2.COLOR_RGB2Lab).astype(np.float32)
    _hsv = cv2.cvtColor(_dig.rgb, cv2.COLOR_RGB2HSV)
    _dig.s = _hsv[:, :, 1].astype(np.float32)
    _dig.v = _hsv[:, :, 2].astype(np.float32)
    _dig.a_map = _dig.lab[:, :, 1] - 128
    _dig.b_map = _dig.lab[:, :, 2] - 128
    _dig.chroma_map = np.hypot(_dig.a_map, _dig.b_map)
    _dig.ang_map = np.degrees(np.arctan2(_dig.b_map, _dig.a_map))
    _dig.assign = None
    _dig.stems = None
    _dig.caps = {}
    _dig.data_points = {}
    _dig._build_ink_mask()
    if _extra_exclude is not None:
        _ex0, _ey0, _ex1, _ey1 = _extra_exclude
        _dig.ink[max(0, _ey0):_ey1 + 1, max(0, _ex0):_ex1 + 1] = False
        _dig.inplot[max(0, _ey0):_ey1 + 1, max(0, _ex0):_ex1 + 1] = False
    recover_palette()
    try:
        with open(os.path.join(OUT_DIR, 'legend_grid.json'), 'w') as _gf:
            json.dump({'legend_box': _lb, 'grid': globals().get('_LAST_LEGEND_GRID'), 'achro': globals().get('_LAST_ACHRO_SWATCHES', []), 'unified': _unified}, _gf)
    except Exception:
        pass
    set_palette_names()
    lock_palette()


def _export_results():
    global Font, PatternFill, _, _CURVE_LABEL, _DBGM, _OUTLIER_ON, _acm, _asg, _bgr, _c, _cal, _canvas, _cd, _cell, _chm, _cix, _cn, _cname, _cnt, _cols, _coords, _curve_ink_mask, _curve_rgb, _cx, _cy, _dbg, _dd, _det, _drp, _dself, _dx, _dy, _e, _ed, _ef, _far, _filter_fake_segments, _filtered, _fname, _fp, _fx_hi, _fx_kind, _fx_lo, _fy_hi, _fy_kind, _fy_lo, _grid, _grid_src, _has_manual, _hdr, _hexcol, _hsv, _hx0, _hx1, _hy0, _hy1, _idx, _ink, _j, _judge, _k, _keep, _key, _label, _lo, _lrgb, _lum, _m, _mask_manifest, _mask_panels, _med_actual, _mf, _mpx, _mr, _msg, _ncols, _nd, _nfar, _nm, _nn, _norm_fallback, _nrows, _ntot, _nx2v, _ny2v, _oc, _oth_names, _ov, _ov_path, _own, _p, _pa_x0, _pa_x1, _pa_y0, _pa_y1, _pal, _pal_map, _panel, _patch, _pax0, _pax1, _pay0, _pay1, _pfx0, _pfx1, _pfy0, _pfy1, _ph, _pts, _pw, _px, _py, _r, _rgb, _rgbf, _rix, _rownum, _rows, _ry, _samp, _sc, _seg_on_curve, _sr, _sw_map, _top, _tot_drop, _txt, _txtcol, _u, _ug, _uni, _user_x, _user_y, _v46_annotate_edit_data, _v46_guard_calibration, _vc, _wb, _ws, _x, _x0, _x1, _x2v, _xhi, _xk, _xkind, _xlo, _xlog, _xlsx_path, _y, _y0, _y1, _y2v, _yhi, _yk, _ykind, _ylo, _ylog, all_detections_out, all_walk_nodes, ax, cd, cname, cov, det, dets, ex, ey, f, fig, get_column_letter, i, json_path, k, label, mean_rgb, n, nodes, nodes_c, openpyxl, p, panel_mask, raw, removed, scale, sx, sy, total, v, x0n, x1n, y0n, y1n
    for cd in COLORS:
        _k = _names.index(cd['name'])
        cd['_raw_mask_canvas'] = (_dig.assign == _k).astype(np.uint8) * 255
        cd['_clean_mask'] = _dig.assign == _k
    _OUTLIER_ON = bool(_args.outlier_filter) and (not os.environ.get('NO_OUTLIER_FILTER'))
    print('  [outlier filter] ' + ('ON (--outlier-filter)' if _OUTLIER_ON else 'OFF (default; step 5 replaces it)'))
    if _OUTLIER_ON:
        _tot_drop = 0
        for _cn in list(all_detections.keys()):
            _filtered, _nd = _filter_line_outliers(all_detections[_cn])
            if _nd > 0:
                all_detections[_cn] = _filtered
                _tot_drop += _nd
                print(f'  [outlier filter] {_cn}: dropped {_nd} off-curve point(s)')

        def _curve_ink_mask(_cn):
            _sw = None
            try:
                for _cd in COLORS:
                    if _cd.get('name') == _cn:
                        _sw = _cd.get('swatch_rgb') or _cd.get('mean_rgb')
                        break
            except NameError:
                _sw = None
            if _sw is None:
                return None
            _swr = np.array([int(_sw[0]), int(_sw[1]), int(_sw[2])], np.int32)
            _tol = 70.0
            try:
                for _cd in COLORS:
                    _os = _cd.get('swatch_rgb') or _cd.get('mean_rgb')
                    if _cd.get('name') == _cn or _os is None:
                        continue
                    _dd = float(np.linalg.norm(_swr - np.array(_os, np.int32)))
                    _tol = min(_tol, max(40.0, _dd * 0.5))
            except NameError:
                pass
            _d = np.linalg.norm(img_rgb.astype(np.int32) - _swr[None, None, :], axis=2)
            _m = (_d < _tol).astype(np.uint8)
            try:
                _px0, _py0, _px1, _py1 = [int(v) for v in PLOT_AREA]
                _band = np.zeros_like(_m)
                _band[_py0:_py1, _px0:_px1] = 1
                _m = _m & _band
            except (NameError, TypeError):
                pass
            if _m.sum() == 0:
                return None
            return cv2.dilate(_m, np.ones((3, 3), np.uint8), iterations=3).astype(bool)

        def _seg_on_curve(p0, p1, ink):
            _n = max(2, int(np.hypot(p1[0] - p0[0], p1[1] - p0[1])) // 2)
            _xs = np.clip(np.linspace(p0[0], p1[0], _n).astype(int), 0, W - 1)
            _ys = np.clip(np.linspace(p0[1], p1[1], _n).astype(int), 0, H - 1)
            return float(ink[_ys, _xs].mean())

        def _filter_fake_segments(pts, ink, seg_thresh=0.12, min_keep=4, max_iter=3):
            if ink is None or len(pts) <= min_keep:
                return (pts, [])
            P = sorted(pts, key=lambda d: d['x'])
            dropped = []
            for _ in range(max_iter):
                n = len(P)
                if n <= min_keep:
                    break
                xy = [(p['x'], p['y']) for p in P]
                svals = [_seg_on_curve(xy[i], xy[i + 1], ink) for i in range(n - 1)]
                worst = -1
                worst_sum = 2.0
                for i in range(1, n - 1):
                    if svals[i - 1] < seg_thresh and svals[i] < seg_thresh:
                        s = svals[i - 1] + svals[i]
                        if s < worst_sum:
                            worst_sum = s
                            worst = i
                if worst < 0:
                    break
                dropped.append((xy[worst], round(svals[worst - 1], 2), round(svals[worst], 2)))
                del P[worst]
            return (P, dropped)
        for _cn in list(all_detections.keys()):
            _ink = _curve_ink_mask(_cn)
            _filtered, _drp = _filter_fake_segments(all_detections[_cn], _ink)
            if _drp:
                all_detections[_cn] = _filtered
                _tot_drop += len(_drp)
                _det = ', '.join((f'({p[0]},{p[1]}) segs={l}/{r}' for p, l, r in _drp))
                print(f'  [fake-segment filter] {_cn}: dropped {len(_drp)} -> {_det}')
        if _tot_drop == 0:
            print('  [outlier filter] no off-curve points found')
    all_detections_out = _V46_MARKERS.output_detections(all_detections)
    removed = [k for k in all_detections if k not in all_detections_out]
    if removed:
        print(f'\n  Post-filter: removed {removed}')
    all_walk_nodes = {}
    for cname, (nodes, cov, *_) in all_results.items():
        if cname in all_detections_out:
            all_walk_nodes[cname] = [[int(p[0]), int(p[1])] for p in nodes]
    json_path = os.path.join(OUT_DIR, 'detections.json')
    with open(json_path, 'w') as f:
        json.dump({'detections': all_detections_out, 'tcaps': {k: all_tcaps_out.get(k, []) for k in all_detections_out}, 'walk_nodes': all_walk_nodes}, f, indent=2)
    print(f'\nJSON: {json_path}')
    try:
        _cal = calibrate_from_axes(img, list(AXIS_ROWS), list(AXIS_COLS), PLOT_AREA)
    except Exception as _e:
        _cal = None
        print(f'  [calibration skipped: {_e}]')
    _x2v = _cal['x2v'] if _cal else None
    _y2v = _cal['y2v'] if _cal else None
    _coords = _cal['coords'] if _cal else None
    _has_manual = any((v is not None for v in (USER_X_MIN, USER_X_MAX, USER_Y_MIN, USER_Y_MAX))) or USER_X_LOG or USER_Y_LOG
    if PLOT_AREA is not None and _has_manual:
        _pa_x0, _pa_y0, _pa_x1, _pa_y1 = PLOT_AREA
        _oc = _coords or {}
        _xlo = USER_X_MIN if USER_X_MIN is not None else _oc.get('x_min')
        _xhi = USER_X_MAX if USER_X_MAX is not None else _oc.get('x_max')
        _xlog = USER_X_LOG or _oc.get('x_kind') == 'log'
        _ylo = USER_Y_MIN if USER_Y_MIN is not None else _oc.get('y_min')
        _yhi = USER_Y_MAX if USER_Y_MAX is not None else _oc.get('y_max')
        _ylog = USER_Y_LOG or _oc.get('y_kind') == 'log'
        _nx2v, _xkind = _make_p2v(_pa_x0, _pa_x1, _xlo, _xhi, _xlog)
        _ny2v, _ykind = _make_p2v(_pa_y1, _pa_y0, _ylo, _yhi, _ylog)
        _user_x = USER_X_MIN is not None or USER_X_MAX is not None or USER_X_LOG
        _user_y = USER_Y_MIN is not None or USER_Y_MAX is not None or USER_Y_LOG
        if _user_x and _nx2v is not None:
            _x2v = _nx2v
        if _user_y and _ny2v is not None:
            _y2v = _ny2v
        if _x2v is not None and _y2v is not None:
            _fx_lo = _xlo if _user_x and _nx2v is not None else _oc.get('x_min')
            _fx_hi = _xhi if _user_x and _nx2v is not None else _oc.get('x_max')
            _fy_lo = _ylo if _user_y and _ny2v is not None else _oc.get('y_min')
            _fy_hi = _yhi if _user_y and _ny2v is not None else _oc.get('y_max')
            _fx_kind = _xkind if _user_x and _nx2v is not None else _oc.get('x_kind', 'linear')
            _fy_kind = _ykind if _user_y and _ny2v is not None else _oc.get('y_kind', 'linear')
            _coords = {'x_min': _fx_lo, 'x_max': _fx_hi, 'y_min': _fy_lo, 'y_max': _fy_hi, 'x_kind': _fx_kind, 'y_kind': _fy_kind}
            print(f'  Calibration: X[{_fx_lo},{_fx_hi}] {_fx_kind}  Y[{_fy_lo},{_fy_hi}] {_fy_kind}  (manual where given)')
        else:
            print('  [manual calibration incomplete; using OCR values where available]')
    _curve_rgb = {}
    _pal_map = {}
    try:
        for _cd in COLORS:
            _nm = _cd.get('name')
            _mr = _cd.get('swatch_rgb') or _cd.get('mean_rgb')
            if _nm is not None and _mr is not None and (len(_mr) == 3):
                _pal_map[_nm] = (int(_mr[0]), int(_mr[1]), int(_mr[2]))
    except Exception:
        _pal_map = {}
    for _cname, _pts in all_detections_out.items():
        if _cname in _pal_map:
            _curve_rgb[_cname] = _pal_map[_cname]
            continue
        _samp = []
        for _p in _pts:
            _x, _y = (int(_p['x']), int(_p['y']))
            _patch = img[max(0, _y - 2):_y + 3, max(0, _x - 2):_x + 3].reshape(-1, 3)
            _samp.extend(_patch.tolist())
        if _samp:
            _m = np.median(_samp, axis=0)
            _curve_rgb[_cname] = (int(_m[2]), int(_m[1]), int(_m[0]))
        else:
            _curve_rgb[_cname] = (0, 0, 0)
    _CURVE_LABEL = _V46_MARKERS.labels(_curve_labels, globals(), _V46_LEGEND)
    from calibration_guard_v46 import guard_calibration as _v46_guard_calibration, annotate_edit_data as _v46_annotate_edit_data
    _v46_guard_calibration(globals())
    _norm_fallback = False
    if (_x2v is None or _y2v is None) and PLOT_AREA is not None:
        _pfx0, _pfy0, _pfx1, _pfy1 = PLOT_AREA
        if _pfx1 != _pfx0 and _pfy1 != _pfy0:
            if _x2v is None:
                _x2v = lambda px, a=_pfx0, b=_pfx1: float((px - a) / (b - a))
            if _y2v is None:
                _y2v = lambda py, a=_pfy1, b=_pfy0: float((py - a) / (b - a))
            _norm_fallback = True
            _coords = {'x_min': 0.0, 'x_max': 1.0, 'y_min': 0.0, 'y_max': 1.0, 'x_kind': 'normalized', 'y_kind': 'normalized'}
            print('  Calibration: none -> normalized [0,1] over plot area')
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
        _wb = openpyxl.Workbook()
        _ws = _wb.active
        _ws.title = 'data_points'
        _hdr = ['curve', 'data_x', 'data_y']
        _ws.append(_hdr)
        for _c in range(1, len(_hdr) + 1):
            _ws.cell(1, _c).font = Font(bold=True, color='FFFFFF')
            _ws.cell(1, _c).fill = PatternFill('solid', start_color='305496')
        _rownum = 1
        for _cname, _pts in all_detections_out.items():
            _label = _CURVE_LABEL.get(_cname, _cname)
            _rgb = _curve_rgb.get(_cname, (0, 0, 0))
            _hexcol = '{:02X}{:02X}{:02X}'.format(max(0, min(255, _rgb[0])), max(0, min(255, _rgb[1])), max(0, min(255, _rgb[2])))
            _lum = 0.299 * _rgb[0] + 0.587 * _rgb[1] + 0.114 * _rgb[2]
            _txtcol = 'FFFFFF' if _lum < 140 else '000000'
            for _p in _pts:
                _px, _py = (_p['x'], _p['y'])
                _dx = _x2v(_px) if _x2v else None
                _dy = _y2v(_py) if _y2v else None
                _ws.append([_label, round(_dx, 4) if _dx is not None else None, round(_dy, 6) if _dy is not None else None])
                _rownum += 1
                _cell = _ws.cell(_rownum, 1)
                _cell.fill = PatternFill('solid', start_color=_hexcol)
                _cell.font = Font(color=_txtcol)
        _ws.append([])
        if _coords and _coords.get('x_kind') == 'normalized':
            _ws.append(['calibration', 'normalized [0,1] over plot area (no axis values)'])
        elif _coords:
            _ws.append(['calibration', f'X[{_coords['x_min']:.4g},{_coords['x_max']:.4g}] {_coords['x_kind']}', f'Y[{_coords['y_min']:.4g},{_coords['y_max']:.4g}] {_coords['y_kind']}'])
        else:
            _ws.append(['calibration', 'not available'])
        _ws.column_dimensions['A'].width = 26
        _ws.column_dimensions['B'].width = 14
        _ws.column_dimensions['C'].width = 14
        _xlsx_path = os.path.join(OUT_DIR, 'data_points.xlsx')
        _wb.save(_xlsx_path)
        print(f'Excel: {_xlsx_path}')
    except Exception as _e:
        print(f'  [Excel skipped: {_e}]')
    try:
        _ed = {'image': {'width': int(W), 'height': int(H)}, 'plot_area': [int(v) for v in PLOT_AREA] if PLOT_AREA else None, 'calibration': None, 'curves': []}
        if PLOT_AREA is not None and _x2v is not None and (_y2v is not None):
            _pax0, _pay0, _pax1, _pay1 = PLOT_AREA
            _xk = (_coords or {}).get('x_kind', 'linear')
            _yk = (_coords or {}).get('y_kind', 'linear')
            _ed['calibration'] = {'x': {'p0': int(_pax0), 'p1': int(_pax1), 'v0': float(_x2v(_pax0)), 'v1': float(_x2v(_pax1)), 'log': _xk == 'log', 'kind': _xk}, 'y': {'p0': int(_pay1), 'p1': int(_pay0), 'v0': float(_y2v(_pay1)), 'v1': float(_y2v(_pay0)), 'log': _yk == 'log', 'kind': _yk}}
        for _cname, _pts in all_detections_out.items():
            _rgb = _curve_rgb.get(_cname, (0, 0, 0))
            _ed['curves'].append({'name': _cname, 'label': _CURVE_LABEL.get(_cname, _cname), 'rgb': [int(_rgb[0]), int(_rgb[1]), int(_rgb[2])], 'points': [{'x': int(_p['x']), 'y': int(_p['y'])} for _p in _pts]})
        with open(os.path.join(OUT_DIR, 'edit_data.json'), 'w') as _ef:
            json.dump(_v46_annotate_edit_data(_ed, globals()), _ef)
        print('  [edit data -> edit_data.json]')
    except Exception as _e:
        print(f'  [edit data skipped: {_e}]')
    try:
        _ov = img.copy()
        for _cname, _pts in all_detections_out.items():
            _rgb = _curve_rgb[_cname]
            _bgr = (_rgb[2], _rgb[1], _rgb[0])
            for _p in _pts:
                _x, _y = (int(_p['x']), int(_p['y']))
                cv2.circle(_ov, (_x, _y), 5, (0, 0, 0), -1)
                cv2.circle(_ov, (_x, _y), 3, _bgr, -1)
                cv2.circle(_ov, (_x, _y), 6, (0, 0, 0), 1)
        _ov_path = os.path.join(OUT_DIR, 'data_points_overlay.png')
        cv2.imwrite(_ov_path, _ov)
        print(f'Overlay: {_ov_path}')
    except Exception as _e:
        print(f'  [Overlay skipped: {_e}]')
    try:
        _asg = getattr(_dig, 'assign', None)
        if _asg is not None:
            _pal = getattr(_dig, 'palette', None)
            _mask_manifest = []
            _DBGM = os.environ.get('DEBUG_MASKS')
            _sw_map = {}
            try:
                for _cd in COLORS:
                    _nm = _cd.get('name')
                    _sr = _cd.get('swatch_rgb') or _cd.get('mean_rgb')
                    if _nm is not None and _sr is not None:
                        _sw_map[_nm] = np.array([int(_sr[0]), int(_sr[1]), int(_sr[2])])
            except Exception:
                _sw_map = {}
            for _k, _cname in enumerate(_names):
                if getattr(_dig, 'is_sink', None) is not None and _k < len(_dig.is_sink) and _dig.is_sink[_k]:
                    continue
                _m = _asg == _k
                if _DBGM and _m.sum() > 0 and (_cname in _sw_map):
                    _own = _sw_map[_cname]
                    _mpx = img_rgb[_m].astype(int)
                    _med_actual = np.median(_mpx, axis=0).astype(int)
                    _dself = np.linalg.norm(_mpx - _own[None, :], axis=1)
                    _far = _dself > 60
                    _nfar = int(_far.sum())
                    _ntot = int(_m.sum())
                    _msg = '  [mask %02d %s] own=%s actual_median=%s  %d px, foreign %d (%d%%)' % (_k, _cname, tuple((int(v) for v in _own)), tuple((int(v) for v in _med_actual)), _ntot, _nfar, round(100 * _nfar / max(_ntot, 1)))
                    if _nfar > 0:
                        _oth_names = [n for n in _sw_map if n != _cname]
                        if _oth_names:
                            _oc = np.stack([_sw_map[n] for n in _oth_names])
                            _fp = _mpx[_far]
                            _dd = np.linalg.norm(_fp[:, None, :] - _oc[None, :, :], axis=2)
                            _nn = _dd.argmin(axis=1)
                            _cnt = {}
                            for _j in _nn:
                                _cnt[_oth_names[_j]] = _cnt.get(_oth_names[_j], 0) + 1
                            _top = sorted(_cnt.items(), key=lambda kv: -kv[1])[:3]
                            _msg += ' | bleed-> ' + ', '.join(('%s:%d' % (n, c) for n, c in _top))
                    print(_msg, flush=True)
                _canvas = np.full((H, W, 3), 255, np.uint8)
                _canvas[_m] = img[_m]
                _pts = all_detections_out.get(_cname, [])
                for _p in _pts:
                    _x, _y = (int(_p['x']), int(_p['y']))
                    cv2.circle(_canvas, (_x, _y), 5, (255, 255, 255), -1)
                    cv2.circle(_canvas, (_x, _y), 5, (0, 0, 0), 2)
                _fname = f'colormask_{_k:02d}_{_cname}.png'
                cv2.imwrite(os.path.join(OUT_DIR, _fname), _canvas)
                _lrgb = [int(v) for v in _pal[_k]] if _pal is not None and _k < len(_pal) else None
                _mask_manifest.append({'name': _cname, 'file': _fname, 'label': _CURVE_LABEL.get(_cname, _cname), 'mask_px': int(_m.sum()), 'points': len(_pts), 'legend_rgb': _lrgb})
            with open(os.path.join(OUT_DIR, 'colormasks.json'), 'w') as _mf:
                json.dump(_mask_manifest, _mf)
            print(f'  [colour masks: {len(_mask_manifest)} -> colormask_*.png]')
            try:
                _ug = globals().get('_LAST_LEGEND_GRID') or {}
                _uni = globals().get('_LAST_UNIFIED_GRID')
            except Exception:
                _uni = None
            try:
                _lo = img.copy()
                if _lb is not None:
                    cv2.rectangle(_lo, (_lb[0], _lb[1]), (_lb[2], _lb[3]), (255, 120, 0), 2)
                _grid_src = _uni if _uni else _ug
                if _grid_src and _grid_src.get('cols') and _grid_src.get('rows'):
                    _cols = _grid_src['cols']
                    _rows = _grid_src['rows']
                    _y0 = _lb[1] if _lb else min(_rows)
                    _y1 = _lb[3] if _lb else max(_rows)
                    _x0 = _lb[0] if _lb else min(_cols)
                    _x1 = _lb[2] if _lb else max(_cols)
                    for _cx in _cols:
                        cv2.line(_lo, (int(_cx), int(_y0)), (int(_cx), int(_y1)), (0, 180, 0), 1)
                    for _ry in _rows:
                        cv2.line(_lo, (int(_x0), int(_ry)), (int(_x1), int(_ry)), (255, 0, 255), 1)
                    if _uni and isinstance(_uni.get('cells'), dict):
                        for _key, _cell in _uni['cells'].items():
                            _cix, _rix = map(int, _key.split(','))
                            _cx, _cy = (int(_cols[_cix]), int(_rows[_rix]))
                            _rgb = _cell.get('rgb', [0, 0, 0])
                            _bgr = (int(_rgb[2]), int(_rgb[1]), int(_rgb[0]))
                            cv2.rectangle(_lo, (_cx - 7, _cy - 7), (_cx + 7, _cy + 7), _bgr, -1)
                            cv2.rectangle(_lo, (_cx - 7, _cy - 7), (_cx + 7, _cy + 7), (0, 0, 0), 1)
                        for _key in _uni.get('empty', []):
                            _cix, _rix = map(int, _key.split(','))
                            _cx, _cy = (int(_cols[_cix]), int(_rows[_rix]))
                            cv2.drawMarker(_lo, (_cx, _cy), (140, 140, 140), cv2.MARKER_TILTED_CROSS, 12, 1)
                cv2.imwrite(os.path.join(OUT_DIR, 'legend_overlay.png'), _lo)
                print('  [legend overlay -> legend_overlay.png]')
            except Exception as _e:
                print(f'  [legend overlay skipped: {_e}]')
            try:
                if _lb is not None:
                    _hx0, _hy0, _hx1, _hy1 = _lb
                    _hsv = cv2.cvtColor(img, cv2.COLOR_BGR2HSV)
                    _sc = _hsv[:, :, 1]
                    _vc = _hsv[:, :, 2]
                    _rgbf = cv2.cvtColor(img, cv2.COLOR_BGR2RGB)
                    _keep = np.zeros(img.shape[:2], np.uint8)
                    _keep[max(0, _hy0):_hy1 + 1, max(0, _hx0):_hx1 + 1] = 1
                    _chm = ((_sc > 40) & (_vc > 40) & (_vc < 250)).astype(np.uint8) & _keep
                    _acm = ((_sc <= 40) & (_vc < 220) & (_vc > 25)).astype(np.uint8) & _keep
                    _dbg = img.copy()
                    cv2.rectangle(_dbg, (_hx0, _hy0), (_hx1, _hy1), (255, 120, 0), 2)

                    def _judge(mask, achro):
                        mask = cv2.morphologyEx(mask, cv2.MORPH_CLOSE, np.ones((3, 3), np.uint8))
                        n, lbl, st, cen = cv2.connectedComponentsWithStats(mask, 8)
                        for i in range(1, n):
                            a = st[i, cv2.CC_STAT_AREA]
                            w = st[i, cv2.CC_STAT_WIDTH]
                            h = st[i, cv2.CC_STAT_HEIGHT]
                            x, y = (st[i, cv2.CC_STAT_LEFT], st[i, cv2.CC_STAT_TOP])
                            cx, cy = cen[i]
                            if not (15 <= a <= 2500 and w <= 120 and (h <= 60)):
                                col = (0, 0, 255)
                            else:
                                tx0, tx1 = (int(cx + 10), int(min(W, cx + 160)))
                                ty0, ty1 = (int(max(0, cy - 12)), int(min(H, cy + 12)))
                                txt = int(((_vc[ty0:ty1, tx0:tx1] < 120) & (_sc[ty0:ty1, tx0:tx1] < 60)).sum()) if tx1 > tx0 else 0
                                rgb = tuple((int(z) for z in np.median(_rgbf[lbl == i], axis=0)))
                                if txt < 100:
                                    col = (0, 220, 255)
                                elif not achro and max(rgb) - min(rgb) <= 24:
                                    col = (0, 150, 255)
                                else:
                                    col = (0, 200, 0) if not achro else (255, 200, 0)
                            cv2.rectangle(_dbg, (x, y), (x + w, y + h), col, 1)
                    _judge(_chm, achro=False)
                    _judge(_acm, achro=True)
                    _u = globals().get('_LAST_UNIFIED_GRID')
                    if _u and _u.get('cols') and _u.get('rows'):
                        for _cx in _u['cols']:
                            cv2.line(_dbg, (int(_cx), _hy0), (int(_cx), _hy1), (0, 180, 0), 1)
                        for _ry in _u['rows']:
                            cv2.line(_dbg, (_hx0, int(_ry)), (_hx1, int(_ry)), (255, 0, 255), 1)
                    _yk = _hy1 + 16
                    for _txt, _c in [('green=chrom OK', (0, 200, 0)), ('cyan=achrom OK', (255, 200, 0)), ('orange=neutral', (0, 150, 255)), ('yellow=no label', (0, 220, 255)), ('red=bad size', (0, 0, 255))]:
                        cv2.putText(_dbg, _txt, (_hx0, _yk), cv2.FONT_HERSHEY_SIMPLEX, 0.4, _c, 1)
                        _yk += 14
                    cv2.imwrite(os.path.join(OUT_DIR, 'legend_diagnostic.png'), _dbg)
                    print('  [legend diagnostic -> legend_diagnostic.png]')
            except Exception as _e:
                print(f'  [legend diagnostic skipped: {_e}]')
    except Exception as _e:
        print(f'  [colour masks skipped: {_e}]')
    if os.environ.get('DEBUG_DUMPS'):
        fig, ax = plt.subplots(figsize=(max(8, W / 60), max(5, H / 60)))
        ax.imshow(img_rgb)
        ax.set_title('Point detection -- all colours', fontsize=11)
        for cname, dets in all_detections_out.items():
            for det in dets:
                ax.plot(det['x'], det['y'], 'o', color='red', ms=5, markeredgecolor='white', markeredgewidth=0.4, zorder=10)
                ax.text(det['x'], det['y'] - 7, f'{det.get('fitness', 0):.2f}', color='yellow', fontsize=4, ha='center', zorder=11)
        ax.axis('off')
        fig.savefig(os.path.join(OUT_DIR, 'points_combined.jpg'), dpi=150, bbox_inches='tight')
        plt.close(fig)
    _mask_panels = []
    for cd in COLORS:
        cname = cd['name']
        if cname not in all_detections_out:
            continue
        mean_rgb = cd.get('mean_rgb', (128, 128, 128))
        raw = cd.get('_raw_mask')
        if raw is None:
            continue
        panel_mask = np.full((H, W, 3), 255, dtype=np.uint8)
        panel_mask[raw == 1] = list(mean_rgb)
        panel_mask = cv2.cvtColor(panel_mask, cv2.COLOR_RGB2BGR)
        nodes_c = all_walk_nodes.get(cname, [])
        if len(nodes_c) > 1:
            for i in range(len(nodes_c) - 1):
                x0n, y0n = nodes_c[i]
                x1n, y1n = nodes_c[i + 1]
                if abs(x1n - x0n) <= 3:
                    cv2.line(panel_mask, (x0n, y0n), (x1n, y1n), (0, 180, 0), 2)
                else:
                    cv2.line(panel_mask, (x0n, y0n), (x1n, y1n), (0, 140, 255), 1)
        if nodes_c:
            sx, sy = nodes_c[0]
            ex, ey = nodes_c[-1]
            cv2.circle(panel_mask, (sx, sy), 9, (0, 255, 0), -1)
            cv2.circle(panel_mask, (sx, sy), 9, (255, 255, 255), 2)
            cv2.putText(panel_mask, 'L', (sx + 11, sy + 5), cv2.FONT_HERSHEY_SIMPLEX, 0.45, (0, 200, 0), 2)
            cv2.circle(panel_mask, (ex, ey), 9, (0, 220, 255), -1)
            cv2.circle(panel_mask, (ex, ey), 9, (255, 255, 255), 2)
            cv2.putText(panel_mask, 'R', (ex + 11, ey + 5), cv2.FONT_HERSHEY_SIMPLEX, 0.45, (0, 200, 220), 2)
        for det in all_detections_out.get(cname, []):
            cv2.circle(panel_mask, (int(det['x']), int(det['y'])), 7, (0, 0, 220), -1)
            cv2.circle(panel_mask, (int(det['x']), int(det['y'])), 7, (255, 255, 255), 1)
        label = f'{cname}  pts={len(all_detections_out.get(cname, []))}  L=green R=cyan'
        cv2.putText(panel_mask, label, (4, 18), cv2.FONT_HERSHEY_SIMPLEX, 0.45, (0, 0, 0), 3)
        cv2.putText(panel_mask, label, (4, 18), cv2.FONT_HERSHEY_SIMPLEX, 0.45, (255, 255, 255), 1)
        scale = 380 / W
        panel_mask = cv2.resize(panel_mask, (380, int(H * scale)), interpolation=cv2.INTER_AREA)
        _mask_panels.append(panel_mask)
    if _mask_panels:
        _ncols = 3
        _nrows = (_ncols + len(_mask_panels) - 1) // _ncols
        _nrows = (len(_mask_panels) + _ncols - 1) // _ncols
        _ph, _pw = _mask_panels[0].shape[:2]
        _grid = np.full((_nrows * _ph + 44, _ncols * _pw, 3), 255, dtype=np.uint8)
        cv2.putText(_grid, 'Raw masks + detected points', (6, 30), cv2.FONT_HERSHEY_SIMPLEX, 0.6, (30, 30, 30), 1)
        for _idx, _panel in enumerate(_mask_panels):
            _r, _c = divmod(_idx, _ncols)
            _grid[44 + _r * _ph:44 + _r * _ph + _ph, _c * _pw:_c * _pw + _pw] = _panel
        cv2.imwrite(os.path.join(OUT_DIR, 'color_masks_montage.jpg'), _grid, [cv2.IMWRITE_JPEG_QUALITY, 93])
    generate_montage(OUT_DIR, img, all_detections_out, all_tcaps_out, all_results)
    if os.environ.get('DEBUG_DUMPS'):
        save_aligned_density_montage(OUT_DIR, COLORS, all_detections_out, all_walk_nodes, all_results, all_wm_full)
    total = sum((len(v) for v in all_detections_out.values()))
    print(f'\nTotal: {total} data points  ->  {OUT_DIR}')
    print('Done.')
