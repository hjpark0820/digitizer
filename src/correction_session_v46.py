"""Portable image + JSON correction sessions; never invoke a point detector."""
from copy import deepcopy
import base64
import hashlib
import io
import json
import math
from pathlib import Path
import re
import zipfile
import zlib

import cv2
import numpy as np

FORMAT = 'chartocode-v46-correction-session-v1'
LINKED_FORMAT = 'chartocode-v46-image-sidecar-v1'
MAX_JSON = 128 * 1024 * 1024
MAX_IMAGE = 64 * 1024 * 1024
MAX_PIXELS = 40_000_000
NPZ = 'color_step5_evidence_v46.npz'
JSON_FILES = ('step5_inputs.json', 'color_correction_state.json', 'correction_state.json',
              'type3_detection_v46.json', 'legend_optional_v46.json', 'step5_stop_report.json',
              'color_group_state_v46.json')
SAFE_ID = re.compile(r'^[A-Za-z0-9_-][A-Za-z0-9_.-]{0,100}$')


def read(path):
    return loads(Path(path).read_bytes())


def loads(data):
    if len(data) > MAX_JSON:
        raise ValueError('Correction JSON exceeds 128 MB')
    def invalid(value):
        raise ValueError('Nonfinite JSON number: '+value)
    value = json.loads(data.decode('utf-8-sig'), parse_constant=invalid)
    check_tree(value)
    return value


def check_tree(value, depth=0):
    if depth > 70:
        raise ValueError('Correction data nesting is too deep')
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError('Nonfinite number in correction data')
    if isinstance(value, dict):
        if 'zlib_base64' in value:
            raw = base64.b64decode(value['zlib_base64'], validate=True)
            obj = zlib.decompressobj()
            decoded = obj.decompress(raw, MAX_PIXELS+1)
            if len(decoded) > MAX_PIXELS or not obj.eof or obj.unused_data:
                raise ValueError('Invalid or oversized compressed reference')
        for v in value.values():
            check_tree(v, depth+1)
    elif isinstance(value, list):
        for v in value:
            check_tree(v, depth+1)


def save(path, value):
    Path(path).write_text(json.dumps(value, ensure_ascii=False, allow_nan=False), encoding='utf-8')


def image_binding(image):
    return dict(width=image.shape[1], height=image.shape[0],
                pixel_sha256=hashlib.sha256(image.tobytes()).hexdigest())


def decode_image(data):
    from PIL import Image
    if len(data) > MAX_IMAGE:
        raise ValueError('Image exceeds 64 MB')
    with Image.open(io.BytesIO(data)) as raster:
        if raster.width*raster.height > MAX_PIXELS:
            raise ValueError('Image exceeds 40 million pixels')
    image = cv2.imdecode(np.frombuffer(data, np.uint8), cv2.IMREAD_COLOR)
    if image is None:
        raise ValueError('Cannot decode the image')
    return image


def validate_edit(ed, image):
    if not isinstance(ed, dict) or not isinstance(ed.get('curves'), list):
        raise ValueError('Expected a correction-session JSON or edit_data.json with curves')
    h, w = image.shape[:2]
    if ed.get('image') != dict(width=w, height=h):
        raise ValueError('Image dimensions do not match the saved data; use the exact paired image')
    if not 1 <= len(ed['curves']) <= 200:
        raise ValueError('Expected 1–200 series')
    plot = ed.get('plot_area')
    if (not isinstance(plot, list) or len(plot) != 4 or
        not all(isinstance(v, (int, float)) and math.isfinite(v) for v in plot) or
        not (0 <= plot[0] < plot[2] < w and 0 <= plot[1] < plot[3] < h)):
        raise ValueError('Invalid saved plot area (inclusive image coordinates required)')
    names = []
    count = 0
    for row in ed['curves']:
        name = row.get('name')
        if not isinstance(name, str) or not SAFE_ID.fullmatch(name):
            raise ValueError('Invalid series identifier')
        names.append(name)
        if not isinstance(row.get('label', name), str):
            raise ValueError('Series label must be text')
        rgb = row.get('rgb')
        if not isinstance(rgb, list) or len(rgb) != 3 or any(
                not isinstance(v, (float,int)) or not math.isfinite(v) or not 0 <= v <= 255 for v in rgb):
            raise ValueError('Invalid series RGB')
        if not isinstance(row.get('points'), list):
            raise ValueError('Series points must be an array')
        count += len(row['points'])
        for p in row['points']:
            if not isinstance(p, dict) or any(not isinstance(p.get(k), (float,int)) or
                not math.isfinite(p[k]) for k in ('x','y')) or not (0 <= p['x'] < w and 0 <= p['y'] < h):
                raise ValueError('Point coordinates must be finite pixels inside the paired image')
    if len(set(names)) != len(names) or count > 100_000:
        raise ValueError('Duplicate series identifiers or too many points')
    if ed.get('calibration') is not None:
        for name in ('x','y'):
            a = ed['calibration'].get(name, {})
            if any(not isinstance(a.get(k), (int,float)) or not math.isfinite(a[k]) for k in ('p0','p1','v0','v1')):
                raise ValueError('Incomplete saved axis calibration')
            if a['p0'] == a['p1'] or (a.get('log') and min(a['v0'],a['v1']) <= 0):
                raise ValueError('Invalid saved axis calibration')
    return deepcopy(ed)


def edited_data(base, proposed, image):
    """Editor may change positions/labels/calibration, never detector evidence."""
    if proposed is None:
        return validate_edit(base, image)
    proposed = validate_edit(proposed, image)
    if {c['name'] for c in base['curves']} != {c['name'] for c in proposed['curves']}:
        raise ValueError('Editing must retain the saved series identities')
    if base['plot_area'] != proposed['plot_area']:
        raise ValueError('Saved plot area cannot change during correction-only editing')
    result = deepcopy(base)
    by_name = {c['name']:c for c in proposed['curves']}
    for row in result['curves']:
        new = by_name[row['name']]
        old = {(p['x'],p['y']):p for p in row['points']}
        row['points'] = [dict(deepcopy(old.get((p['x'],p['y']),{})),x=p['x'],y=p['y']) for p in new['points']]
        row['label'] = new.get('label', row.get('label',row['name']))
    result['calibration'] = proposed.get('calibration')
    return result


def checked_npz(data):
    # Inspect headers BEFORE NumPy allocation; no ZIP extraction or pickle.
    if len(data) > MAX_JSON:
        raise ValueError('Evidence archive is too large')
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        members = archive.infolist()
        if len(members) > 500 or sum(m.file_size for m in members) > 256*1024*1024:
            raise ValueError('Evidence archive expands beyond allowed size')
        if len({m.filename for m in members}) != len(members):
            raise ValueError('Duplicate evidence array')
        for m in members:
            if '/' in m.filename or '\\' in m.filename or not m.filename.endswith('.npy'):
                raise ValueError('Evidence archive must contain sibling numeric arrays only')
            with archive.open(m) as f:
                version = np.lib.format.read_magic(f)
                if version == (1,0):
                    shape, _, dtype = np.lib.format.read_array_header_1_0(f)
                elif version == (2,0):
                    shape, _, dtype = np.lib.format.read_array_header_2_0(f)
                else:
                    raise ValueError('Unsupported evidence array header')
                if dtype.hasobject or len(shape) != 2 or math.prod(shape)*dtype.itemsize > m.file_size:
                    raise ValueError('Invalid numeric evidence array')
    return data


def bw_context(state, image):
    if state.get('format_version') == 'bw_correction_state_v46_v1':
        bind = state['binding']
        if bind.get('upscale') != 1. or bind.get('point_backend') != 'grid_v46':
            raise ValueError('Correction-only requires native-scale v46 B&W data')
        plot, legend = bind['plot_area'], bind.get('legend_area')
        runtime_key = 'bw_step5_state'
    elif state.get('backend') == 'v46_legend_optional':
        bind = state['identity']
        plot, legend = bind['plot_box'], bind.get('legend_box')
        runtime_key = 'runtime_state'
    else:
        raise ValueError('Unsupported B&W correction state')
    if bind.get('image_sha256') != image_binding(image)['pixel_sha256']:
        raise ValueError('B&W evidence belongs to different image pixels')
    from bw_step5_v46 import _box, _normalise
    _box(plot,image.shape); _box(legend,image.shape)
    diameter = float((state.get(runtime_key) or {}).get('diameter',15.))
    _normalise(state['P_current'],state['S_current'],plot,legend,diameter)
    return dict(plot=plot,legend=legend,runtime_key=runtime_key,diameter=diameter)


def capability(folder, image, ed):
    """Missing evidence is manual-only; corrupt evidence is an error."""
    folder = Path(folder)
    if ed.get('correction_available') is False:
        return None, ed.get('correction_unavailable_reason') or 'Saved run does not support automatic Step 5.'
    from color_group_runtime_v46 import STATE_FILE, VERSION as GROUP_VERSION, validate as validate_groups
    if ed.get('detector_backend') == GROUP_VERSION or (folder/STATE_FILE).exists():
        validate_groups(read(folder/STATE_FILE),image,ed)
        return 'color_group_bw', None
    if (folder/'type3_detection_v46.json').exists():
        from type3_v46.gui_adapter import validate_detection
        validate_detection(read(folder/'type3_detection_v46.json'),image,ed)
        return 'type3', None
    if (folder/'step5_inputs.json').exists():
        from color_step5_payload_v46 import load_payload
        raw = read(folder/'step5_inputs.json')
        if raw.get('masks_npz') != NPZ:
            raise ValueError('Unexpected Step-5 evidence archive name')
        checked_npz((folder/NPZ).read_bytes())
        loaded = load_payload(folder,image,raw['plot_area'],raw.get('legend_box'),[c['name'] for c in ed['curves']])
        if loaded is None:
            raise ValueError('Unsupported colour Step-5 evidence version')
        if list(raw['plot_area']) != ed['plot_area']:
            raise ValueError('Edited plot area differs from correction evidence')
        return 'color', None
    if (folder/'correction_state.json').exists():
        state = read(folder/'correction_state.json')
        bw_context(state,image)
        known = {p['swatch_id'] for p in state['P_current']+state['S_current']}
        known.update(c['swatch_id'] for c in state.get('series_catalog',[]) if c.get('swatch_id'))
        if any(c.get('swatch_id',c['name']) not in known for c in ed['curves']):
            raise ValueError('B&W series does not match saved marker evidence')
        return 'bw', None
    return None, 'Coordinates loaded for manual editing. Automatic Step 5 needs a correction-session JSON with saved evidence; detection was not run.'


def export_session(folder, mode, proposed=None):
    folder = Path(folder)
    image = cv2.imread(str(folder/'input.png'))
    if image is None:
        raise ValueError('Paired job image is unavailable')
    ed = edited_data(read(folder/'edit_data.json'), proposed, image)
    artifacts = {name:dict(encoding='json',data=read(folder/name)) for name in JSON_FILES if (folder/name).exists()}
    if (folder/NPZ).exists():
        artifacts[NPZ] = dict(encoding='base64',data=base64.b64encode(checked_npz((folder/NPZ).read_bytes())).decode('ascii'))
    return dict(format=FORMAT, mode=mode, image=image_binding(image), edit_data=ed,
                artifacts=artifacts, coordinate_system='paired_image_pixels',
                note='Keep this JSON with the exact image. Saved hypotheses are not validated measurements.')


def excel_bytes(ed):
    """Application export, using the project's existing openpyxl dependency.

    Always derive values from current pixels, never stale detector data_x/data_y.
    This is a data table, not a substitute for the evidence-bearing session JSON.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    def value(pixel, axis):
        if axis is None or axis.get('kind') == 'normalized':
            return None
        t = (pixel-axis['p0'])/(axis['p1']-axis['p0'])
        if axis.get('log'):
            result = 10**(math.log10(axis['v0'])+t*(math.log10(axis['v1'])-math.log10(axis['v0'])))
        else:
            result = axis['v0']+t*(axis['v1']-axis['v0'])
        if not math.isfinite(result):
            raise ValueError('Axis calibration produces nonfinite values')
        return result

    wb = Workbook()
    ws = wb.active
    ws.title = 'data_points'
    ws.append(['series_id', 'series_label', 'point_index', 'pixel_x', 'pixel_y', 'data_x', 'data_y'])
    cal = ed.get('calibration') or {}
    for curve in ed['curves']:
        for index, point in enumerate(sorted(curve['points'], key=lambda p:(p['x'],p['y'])), 1):
            ws.append([curve['name'], curve.get('label',curve['name']), index, point['x'], point['y'],
                       value(point['x'],cal.get('x')), value(point['y'],cal.get('y'))])
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = 'C2'
    meta = wb.create_sheet('calibration')
    meta.append(['field', 'value'])
    meta.append(['points', 'Active points only. Suppressed hypotheses are retained in the correction JSON.'])
    meta.append(['pixel_coordinates', 'Paired image pixels: origin at top left, x rightward, y downward.'])
    meta.append(['data_coordinates', 'Axis-calibrated values; blank when calibration is unavailable or normalized only.'])
    meta.append(['image_width', ed['image']['width']])
    meta.append(['image_height', ed['image']['height']])
    meta.append(['plot_area', ', '.join(map(str,ed['plot_area']))])
    for name in ('x','y'):
        axis = cal.get(name)
        meta.append([name+'_scale', (axis.get('kind') or ('log10' if axis.get('log') else 'linear')) if axis else 'unavailable'])
        for key in ('p0','p1','v0','v1'):
            meta.append([name+'_'+key, axis[key] if axis else None])
    for sheet, widths in ((ws,[20,48,16,18,18,20,20]),(meta,[24,100])):
        sheet.sheet_view.showGridLines = False
        for row in sheet:
            for cell in row:
                # Treat imported labels as literal strings, including '=...'.
                if isinstance(cell.value,str):
                    cell.data_type = 's'
                cell.font = Font(name='Arial',size=11)
                cell.alignment = Alignment(vertical='center')
        for cell in sheet[1]:
            cell.font = Font(name='Arial',size=11,bold=True,color='FFFFFF')
            cell.fill = PatternFill('solid',fgColor='305496')
            cell.alignment = Alignment(horizontal='center',vertical='center')
        for col,width in enumerate(widths,1):
            sheet.column_dimensions[get_column_letter(col)].width = width
        for row in range(1,sheet.max_row+1):
            sheet.row_dimensions[row].height = 20
    for row in ws.iter_rows(min_row=2,min_col=3,max_col=3):
        row[0].number_format = '0'
    for row in ws.iter_rows(min_row=2,min_col=4,max_col=7):
        for cell in row:
            cell.number_format = '0.000' if cell.column <= 5 else '0.000000E+00'
    for row in ws.iter_rows(min_row=2,min_col=1,max_col=2):
        for cell in row:
            cell.alignment = Alignment(vertical='center',wrap_text=True)
        lines = max(math.ceil(len(str(row[0].value))/19),math.ceil(len(str(row[1].value))/46))
        ws.row_dimensions[row[0].row].height = max(20,16*lines)
    for row in meta.iter_rows(min_row=2):
        row[1].alignment = Alignment(vertical='center',wrap_text=True)
        if isinstance(row[1].value,str) and len(row[1].value)>85:
            meta.row_dimensions[row[1].row].height = 32
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def save_outputs(folder, mode):
    """Snapshot both formats on detection/import/correction completion."""
    folder = Path(folder)
    package = export_session(folder,mode)
    workbook = excel_bytes(package['edit_data'])
    save(folder/'correction_session.json',package)
    (folder/'current_data.xlsx').write_bytes(workbook)


def export_linked_session(folder, mode, original_bytes, filename, proposed=None):
    """Sidecar pairs with the source image, preserving prepared pixels internally."""
    source = decode_image(original_bytes)
    package = export_session(folder, mode, proposed)
    package['format'] = LINKED_FORMAT
    package['source_image'] = dict(image_binding(source), filename=Path(filename).name,
                                  file_sha256=hashlib.sha256(original_bytes).hexdigest())
    if image_binding(source) != package['image']:
        working = cv2.imread(str(Path(folder)/'input.png'))
        ok, png = cv2.imencode('.png',working)
        if not ok:
            raise ValueError('Cannot retain the prepared image')
        package['working_image_png_base64'] = base64.b64encode(png).decode('ascii')
    package['note'] = 'Open the original image with its same-name JSON. Prepared pixels, if any, are stored in this JSON; the original image is never overwritten.'
    if len(json.dumps(package,ensure_ascii=False).encode('utf-8')) > MAX_JSON:
        raise ValueError('Image plus correction evidence exceeds the 128 MB JSON limit')
    return package


def import_session(image_bytes, data_bytes, folder, mode='color', version_id=None):
    image = decode_image(image_bytes)
    value = loads(data_bytes)
    if not isinstance(value,dict):
        raise ValueError('Data file must contain a JSON object')
    from correction_history_v46 import FORMAT as HISTORY_FORMAT, unpack
    history, version = None, None
    if value.get('format') == HISTORY_FORMAT:
        history = value
        value, version = unpack(history, version_id)
    elif version_id:
        raise ValueError('This legacy JSON does not contain selectable versions')
    paired = value.get('format') in (FORMAT, LINKED_FORMAT)
    if 'format' in value and not paired:
        raise ValueError('Unsupported correction-session version')
    if paired:
        if value['format'] == LINKED_FORMAT:
            source = value.get('source_image') or {}
            if any(source.get(k) != v for k,v in image_binding(image).items()):
                raise ValueError('Original image does not match its JSON sidecar; pixels changed')
            if 'working_image_png_base64' in value:
                image = decode_image(base64.b64decode(value['working_image_png_base64'],validate=True))
        if value.get('image') != image_binding(image):
            raise ValueError('Image/data pair mismatch: pixels differ. Use the matching saved PNG without resizing or rotation.')
        mode = value.get('mode')
        ed = validate_edit(value.get('edit_data'),image)
        artifacts = value.get('artifacts',{})
        if not isinstance(artifacts,dict) or set(artifacts)-set(JSON_FILES+(NPZ,)):
            raise ValueError('Unexpected artifact in correction session')
    else:
        ed = validate_edit(value,image)
        artifacts = {}
    if mode not in ('bw','color'):
        raise ValueError('Unknown saved mode')
    folder = Path(folder)
    folder.mkdir(parents=True,exist_ok=False)
    for name, item in artifacts.items():
        if name == NPZ:
            if item.get('encoding') != 'base64':
                raise ValueError('Expected base64 numeric evidence')
            (folder/name).write_bytes(checked_npz(base64.b64decode(item['data'],validate=True)))
        else:
            if item.get('encoding') != 'json' or not isinstance(item.get('data'),dict):
                raise ValueError('Expected JSON evidence object')
            save(folder/name,item['data'])
    engine, reason = capability(folder,image,ed)
    if not cv2.imwrite(str(folder/'input.png'),image):
        raise ValueError('Cannot save paired image')
    save(folder/'edit_data.json',ed)
    meta = dict(mode=mode,engine=engine,imported=True,paired_hash_checked=paired,
                manual_only_reason=reason, detection_skipped=True)
    if history is not None:
        save(folder/'imported_history.json',history)
        meta.update(version_id=version['id'], version_kind=version.get('kind','imported'),
                    created_at=version.get('created_at'), iterations=version.get('iterations'))
    save(folder/'session_meta.json',meta)
    draw_overlay(folder,image,ed)
    return dict(mode=mode,engine=engine,reason=reason,paired=paired,version=version)


def draw_overlay(folder,image,ed):
    from type3_v46.gui_adapter import _overlay
    if not cv2.imwrite(str(Path(folder)/'data_points_overlay.png'),_overlay(image,ed['curves'])):
        raise ValueError('Cannot save edited point overlay')


def sync_bw_series(ed, points, catalog=()):
    """Export every typed active point, including newly activated identities.

    Existing labels, colours, calibration and ordering are never replaced.
    Legacy portable states can recover identity from typed point metadata even
    when they predate the full point-free legend catalog.
    """
    from bw_series_identity import swatch_curve
    rows={c.get('swatch_id',c['name']):c for c in ed['curves']}
    for entry in catalog:
        sid=entry.get('swatch_id',entry['name'])
        if sid not in rows:
            row=deepcopy(entry);row['points']=[]
            ed['curves'].append(row);rows[sid]=row
    for point in points:
        sid=point['swatch_id']
        if sid not in rows:
            row=swatch_curve(point)
            ed['curves'].append(row);rows[sid]=row
    for row in rows.values():
        row['points']=[]
        row.pop('reconstruction',None)
    for point in points:
        rows[point['swatch_id']]['points'].append(dict(x=point['cx'],y=point['cy']))
    for row in rows.values():
        row['points'].sort(key=lambda p:p['x'])
    if sum(len(c['points']) for c in ed['curves']) != len(points):
        raise ValueError('B&W export must preserve every corrected active point exactly once')


def run_bw_saved(folder,image,ed,iterations):
    """Route typed BW series by source connectivity, without run_detection()."""
    from bw_series_correction_v46 import run_correction
    folder = Path(folder)
    state = read(folder/'correction_state.json')
    context = bw_context(state,image)
    templates = {}
    for p in state['P_current']+state['S_current']:
        templates.setdefault(p['swatch_id'],p)
    active = []
    for c in ed['curves']:
        sid = c.get('swatch_id',c['name'])
        old = [p for p in state['P_current'] if p['swatch_id']==sid]
        for i,p in enumerate(c['points']):
            exact = next((q for q in old if abs(q['cx']-p['x'])<1e-6 and abs(q['cy']-p['y'])<1e-6),None)
            if exact is None and sid not in templates:
                raise ValueError(f'B&W series {sid} has no saved marker evidence for automatic correction')
            q = deepcopy(exact or templates[sid])
            q.update(cx=p['x'],cy=p['y'])
            if exact is None:
                q.update(point_id=f'{sid}_manual_{i}',candidate_id=f'{sid}_manual_{i}',
                         original_detection=False,manual_edit=True,confidence=0.)
            active.append(q)
    occupied = {(p['swatch_id'],p['cx'],p['cy']) for p in active}
    pool = [p for p in state['S_current'] if (p['swatch_id'],p['cx'],p['cy']) not in occupied]
    previous = state.get(context['runtime_key'])
    result = run_correction(image,folder/'correction_diagnostics',init_points=active,init_suppressed=pool,
        plot_area=context['plot'],legend_box=context['legend'],d_est=context['diameter'],max_iter=iterations,
        connection_alignment=True,structural_edits=True,
        marker_evidence=state.get('marker_evidence'),
        init_state=previous,no_legend=bool((previous or {}).get('binding',{}).get('explicit_no_legend') or
            (state.get('backend')=='v46_legend_optional' and context['legend'] is None)))
    state.update(P_current=result['P_current'],S_current=result['S_current'])
    state[context['runtime_key']] = result['runtime_state']
    save(folder/'correction_state.json',state)
    sync_bw_series(ed,result['P_current'],state.get('series_catalog',[]))
    ed['correction'] = dict(backend='bw_series_correction_v46',metric='structure_routed',detection_skipped=True,
                            connection_alignment=True,structural_edits=True,
                            series_structure=result['runtime_state'].get('structure',{}))
    save(folder/'edit_data.json',ed)
    draw_overlay(folder,image,ed)


def correct_saved(folder, iterations, stop_policy='manual'):
    from step5_stop_v46 import resolve_policy, trace_snapshot, build_report
    config = resolve_policy(stop_policy, iterations)
    iterations = config['max_iterations']
    folder = Path(folder)
    image = cv2.imread(str(folder/'input.png'))
    if image is None:
        raise ValueError('Cannot read paired image')
    ed = validate_edit(read(folder/'edit_data.json'),image)
    engine, reason = capability(folder,image,ed)
    if engine is None:
        raise ValueError(reason)
    before = trace_snapshot(folder)
    if engine == 'bw':
        run_bw_saved(folder,image,ed,iterations)
    elif engine == 'color_group_bw':
        from color_group_runtime_v46 import correct_saved as correct_groups
        correct_groups(folder,image,ed,iterations)
    else:
        from color_correct_cli import main
        args = [str(folder/'input.png'),str(folder),'--require-v46-path','--use-edit-points',
                '--correct-iters',str(iterations)]
        previous = folder/'color_correction_state.json'
        if previous.exists():
            args += ['--prev-state',str(previous)]
        main(args)
    stop = build_report(folder,config,before)
    save(folder/'correction_only_status.json',dict(completed=True,engine=engine,detection_skipped=True,
        iterations_requested=iterations,stop_policy=config['policy'],stop_report=stop))
